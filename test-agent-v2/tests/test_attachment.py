"""Jira/Confluence attachment reading — the fetcher + extractor that pull PDF/Office/text/image
attachments (e.g. the spec PDFs) into the pack instead of leaving them recorded-only."""

from __future__ import annotations

import io

import pytest

from common.extract.assemble import extract_issue_links
from common.extract.attachment import attachment_text, is_extractable_attachment
from common.extract.classify import classify_url
from common.models import ATTACHMENT, Scope

# --- type gate -------------------------------------------------------------------------------

def test_is_extractable_covers_docs_and_images_not_blobs():
    assert is_extractable_attachment("application/pdf", "spec.pdf")
    assert is_extractable_attachment(None, "NOTES.PDF")                 # by extension
    assert is_extractable_attachment("text/csv", "data.csv")
    assert is_extractable_attachment(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "d.docx")
    assert is_extractable_attachment(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "d.xlsx")
    assert is_extractable_attachment("image/png", "shot.png")
    assert not is_extractable_attachment("application/zip", "a.zip")    # blobs never become nodes
    assert not is_extractable_attachment("video/mp4", "clip.mp4")


# --- routing + scope -------------------------------------------------------------------------

def test_classify_attachment_canonicalizes_to_attachment_scheme():
    url = "https://x.atlassian.net/rest/api/3/attachment/content/12345"
    typ, canon = classify_url(url)
    assert typ == ATTACHMENT
    assert canon == f"attachment:{url}"  # routes to the authed AttachmentFetcher, not the web fetcher


def test_scope_follows_attachments_by_default():
    assert Scope().follows(ATTACHMENT)


def test_assemble_keeps_extractable_attachments_drops_blobs():
    issue = {"key": "LUZ-1", "fields": {"attachment": [
        {"content": "https://x.atlassian.net/rest/api/3/attachment/content/1",
         "filename": "spec.pdf", "mimeType": "application/pdf"},
        {"content": "https://x.atlassian.net/rest/api/3/attachment/content/2",
         "filename": "demo.mp4", "mimeType": "video/mp4"},
    ]}}
    links = extract_issue_links(issue, base_url="https://x.atlassian.net", scope=Scope())
    canons = {lr.canonical_url for lr in links}
    assert any(c.endswith("content/1") for c in canons)      # pdf kept
    assert not any(c.endswith("content/2") for c in canons)  # video filtered at the source
    pdf = next(lr for lr in links if lr.canonical_url.endswith("content/1"))
    assert pdf.type == ATTACHMENT and pdf.in_scope


# --- extraction ------------------------------------------------------------------------------

def test_text_and_csv_extraction():
    assert attachment_text(b"a,b\n1,2", "text/csv", "d.csv") == "a,b\n1,2"


def test_docx_extraction_reads_paragraphs_and_tables():
    from docx import Document

    doc = Document()
    doc.add_paragraph("Hello spec")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "key", "value"
    buf = io.BytesIO()
    doc.save(buf)
    out = attachment_text(
        buf.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "d.docx")
    assert "Hello spec" in out and "key | value" in out


def test_xlsx_extraction_reads_cells():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["Name", "Qty"])
    ws.append(["widget", 3])
    buf = io.BytesIO()
    wb.save(buf)
    out = attachment_text(
        buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "d.xlsx")
    assert "Name" in out and "widget" in out and "3" in out


def test_pdf_branch_runs_without_raising():
    from pypdf import PdfWriter

    w = PdfWriter()
    w.add_blank_page(width=200, height=200)
    buf = io.BytesIO()
    w.write(buf)
    # a blank page has no text, but the pdf branch must execute + return a str (proves pypdf wiring)
    assert isinstance(attachment_text(buf.getvalue(), "application/pdf", "x.pdf"), str)


def test_image_without_vertex_degrades_to_placeholder(monkeypatch):
    for k in ("VERTEX_PROJECT", "VERTEX_LOCATION", "VERTEX_MODEL"):
        monkeypatch.delenv(k, raising=False)
    out = attachment_text(b"\x89PNG\r\n\x1a\n", "image/png", "shot.png")
    assert "image attachment" in out and "shot.png" in out  # recorded, never raises, no fabricated text


def test_oversized_image_is_capped_before_decode(monkeypatch):
    """INT-07: a raw image blob over the decode ceiling degrades to a placeholder — never PIL-decoded."""
    from common.extract.attachment import _MAX_DECODE_BYTES

    for k in ("VERTEX_PROJECT", "VERTEX_LOCATION", "VERTEX_MODEL"):
        monkeypatch.delenv(k, raising=False)
    out = attachment_text(b"\x00" * (_MAX_DECODE_BYTES + 1), "image/png", "huge.png")
    assert "too large" in out and "huge.png" in out  # bounded before Image.open, never raises


def test_unsupported_type_raises_so_crawl_records_a_gap():
    with pytest.raises(ValueError):
        attachment_text(b"PK\x03\x04", "application/zip", "a.zip")


# --- the fetcher (authed download → note) ----------------------------------------------------

async def test_attachment_fetcher_downloads_and_extracts():
    from knowledge_gathering.gather.crawl.fetch import fetch_node

    class FakeClient:
        async def download_bytes(self, url, *, max_bytes=25 * 1024 * 1024):
            assert url == "https://x.atlassian.net/rest/api/3/attachment/content/9"
            return b"line1\nline2", "text/plain", "notes.txt"

    nid = "attachment:https://x.atlassian.net/rest/api/3/attachment/content/9"
    links, note, text = await fetch_node(FakeClient(), nid, Scope())
    assert links == []
    assert note.type == ATTACHMENT and note.title == "notes.txt" and note.id == nid
    assert text == "line1\nline2"
